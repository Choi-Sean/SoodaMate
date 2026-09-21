"""Builds the Excel QA report from qa/out/*.json (+ pytest junit xml).
Run with the *global* python (has openpyxl):  python backend/qa/build_report.py"""
import json
import sys
import xml.etree.ElementTree as ET
from collections import OrderedDict, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
OUT = Path(__file__).resolve().parent / "out"
DEST = Path(__file__).resolve().parents[2] / "docs" / "qa"
DEST.mkdir(parents=True, exist_ok=True)

KR_CAT = {
    "Auth": "인증/가입", "Profile": "프로필", "Upload": "업로드(R2 실제)", "Verification": "인증(얼굴·전화·재직)",
    "Safety": "안전(차단·신고)", "Admin": "관리자", "Discovery": "탐색(디스커버)", "Matching": "스와이프·매칭",
    "Chat": "채팅·통화", "Blind chat": "블라인드 채팅·AI매칭", "Moments": "모먼트", "Couple stories": "커플 스토리",
    "Payments": "결제", "Push": "푸시·기기", "Account": "계정", "Security": "보안", "Ads": "광고", "i18n": "다국어(i18n)",
    "Website": "웹사이트", "Website vs App": "웹↔앱 일치", "UI: App": "화면(앱)", "UI: Admin": "화면(관리자)",
    "UI: Website": "화면(웹사이트)", "Regression (pytest)": "기존 자동회귀(pytest)",
}


def load(name):
    p = OUT / name
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {"results": [], "defects": []}


def by_id(rows):
    d = OrderedDict()
    for r in rows:
        d[r["id"]] = r
    return d


final = OrderedDict()
for f in ("results_core.json", "results_core_only.json", "results_flow.json", "results_static.json", "results_flow_only.json", "results_security2.json", "results_security2_only.json"):
    final.update(by_id(load(f)["results"]))

baseline = OrderedDict()
for f in ("baseline_core.json", "baseline_flow.json", "baseline_security.json"):
    baseline.update(by_id(load(f)["results"]))
# static baseline = first run (before the website/claim fixes)
STATIC_BASE = {"CLAIM-001": 50.0, "CLAIM-002": 0.0, "CLAIM-003": 0.0, "CLAIM-005": 50.0}

from qa.qa_defects import DEFECTS, CLAIMS, LIMITS  # noqa: E402
from qa.ui_results import UI  # noqa: E402

# ---------------------------------------------------------------- pytest xml
pytest_rows = []
junit = OUT / "pytest_junit.xml"
if junit.exists():
    root = ET.parse(junit).getroot()
    for tc in root.iter("testcase"):
        outcome = "PASS"
        detail = ""
        for child in tc:
            if child.tag in ("failure", "error"):
                outcome, detail = "FAIL", (child.get("message") or "")[:300]
            elif child.tag == "skipped":
                outcome, detail = "SKIP", (child.get("message") or "")[:200]
        pytest_rows.append((tc.get("classname", "").replace("tests.", ""), tc.get("name"), outcome, float(tc.get("time", 0)), detail))

# ------------------------------------------------------------------- styles
HDR = PatternFill("solid", fgColor="FF6B9D")
HDR_FONT = Font(bold=True, color="FFFFFF", name="Malgun Gothic")
GREEN = PatternFill("solid", fgColor="D5F5E3")
RED = PatternFill("solid", fgColor="FADBD8")
YELLOW = PatternFill("solid", fgColor="FCF3CF")
GREY = PatternFill("solid", fgColor="EAECEE")
thin = Side(style="thin", color="D0D3D4")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
BASE_FONT = Font(name="Malgun Gothic", size=10)


def sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill, c.font, c.alignment, c.border = HDR, HDR_FONT, Alignment(wrap_text=True, vertical="center"), BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    return ws


def style_rows(ws, result_col=None):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment, c.border, c.font = WRAP, BORDER, BASE_FONT
        if result_col:
            v = str(row[result_col - 1].value or "")
            fill = GREEN if v.startswith("PASS") or v.startswith("Fixed") or v.startswith("수정") else RED if v.startswith("FAIL") or v.startswith("Open") or v.startswith("미해결") else YELLOW if v else None
            if fill:
                row[result_col - 1].fill = fill
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()
wb.remove(wb.active)

# ------------------------------------------------------------ test cases sheet
ws = sheet(wb, "테스트케이스", [
    "케이스 ID", "분류", "기능", "시나리오", "실행 절차", "예상 결과 (Expected)", "실제 결과 (Actual, 수정 후 재실행)",
    "결과", "예상 일치율(%)", "1차 실행(수정 전) 결과", "1차 일치율(%)", "증거/메모", "소요(초)"],
    [11, 16, 22, 48, 34, 52, 60, 9, 11, 13, 11, 40, 9])

all_cases = list(final.values())
for r in all_cases:
    b = baseline.get(r["id"])
    base_res, base_pct = ("-", None)
    if b:
        base_res, base_pct = b["result"], b["match_pct"]
    elif r["id"] in STATIC_BASE:
        base_res, base_pct = ("FAIL" if STATIC_BASE[r["id"]] < 100 else "PASS"), STATIC_BASE[r["id"]]
    ws.append([r["id"], KR_CAT.get(r["category"], r["category"]), r["feature"], r["scenario"], r.get("steps", ""), r["expected"],
               r["actual"][:1500], r["result"], r["match_pct"], base_res, base_pct, r.get("evidence", ""), r.get("seconds")])
for cid, cat, feat, scen, exp, obs, ok, pct in UI:
    all_cases.append({"id": cid, "category": cat, "feature": feat, "scenario": scen, "expected": exp, "actual": obs,
                      "result": "PASS" if ok else ("N/T" if ok is None else "FAIL"), "match_pct": pct, "checks_total": 1 if pct is not None else 0,
                      "checks_passed": (pct or 0) / 100.0})
    ws.append([cid, KR_CAT.get(cat, cat), feat, scen, "브라우저에서 직접 조작", exp, obs, "PASS" if ok else ("N/T" if ok is None else "FAIL"), pct, "-", None, "", None])
style_rows(ws, result_col=8)

# --------------------------------------------------------------- summary sheet
def is_exec(r):
    return r["result"] in ("PASS", "FAIL")


executed = [r for r in all_cases if is_exec(r)]
n_pass = sum(1 for r in executed if r["result"] == "PASS")
n_fail = sum(1 for r in executed if r["result"] == "FAIL")
n_nt = sum(1 for r in all_cases if r["result"] == "N/T")
chk_total = sum(r.get("checks_total") or 0 for r in executed)
chk_pass = sum(r.get("checks_passed") or 0 for r in executed)
avg_match = sum((r["match_pct"] or 0) for r in executed) / max(len(executed), 1)

base_all = [b for b in baseline.values() if b["result"] in ("PASS", "FAIL")]
base_pass = sum(1 for b in base_all if b["result"] == "PASS")
base_avg = sum((b["match_pct"] or 0) for b in base_all) / max(len(base_all), 1)

ws = wb.create_sheet("요약", 0)
ws.column_dimensions["A"].width = 34
for col in "BCDEFGH":
    ws.column_dimensions[col].width = 16
ws["A1"] = "SooDaMate 전수 QA 리포트"
ws["A1"].font = Font(bold=True, size=16, color="FF6B9D", name="Malgun Gothic")
ws["A2"] = f"실행일 {datetime.now():%Y-%m-%d} · 실제 호스팅 DB / R2 스토리지 / Twilio Lookup / Google 번역 · 푸시·SMS·메일 발송은 스텁(실제 사용자 영향 없음)"
ws["A2"].font = BASE_FONT
rows = [
    ("총 케이스(직접 실행)", len(executed)), ("PASS (예상과 100% 일치)", n_pass), ("FAIL", n_fail), ("실행 불가(N/T)", n_nt),
    ("케이스 단위 통과율", f"{100.0 * n_pass / max(len(executed), 1):.1f}%"),
    ("체크 항목 단위 예상 일치율", f"{100.0 * chk_pass / max(chk_total, 1):.1f}%  ({int(chk_pass)}/{chk_total} 체크)"),
    ("케이스 평균 예상 일치율", f"{avg_match:.1f}%"),
    ("", ""),
    ("[참고] 수정 전 1차 실행 통과율", f"{100.0 * base_pass / max(len(base_all), 1):.1f}%  ({base_pass}/{len(base_all)})"),
    ("[참고] 수정 전 1차 평균 일치율", f"{base_avg:.1f}%"),
    ("발견된 결함(수정 완료/열림/정보)", f"{sum(1 for d in DEFECTS if d[3].startswith('수정'))} / {sum(1 for d in DEFECTS if d[3].startswith('미해결'))} / {sum(1 for d in DEFECTS if d[3].startswith('정보'))}"),
    ("기존 pytest 회귀", f"{sum(1 for p in pytest_rows if p[2] == 'PASS')} / {len(pytest_rows)} 통과" if pytest_rows else "미실행"),
]
r0 = 4
for i, (k, v) in enumerate(rows):
    ws.cell(r0 + i, 1, k).font = Font(bold=True, name="Malgun Gothic")
    ws.cell(r0 + i, 2, v).font = BASE_FONT
ws.cell(r0 + len(rows) + 1, 1, "분류별 결과").font = Font(bold=True, size=12, name="Malgun Gothic")
hdr_row = r0 + len(rows) + 2
for j, h in enumerate(["분류", "케이스 수", "PASS", "FAIL", "N/T", "평균 일치율(%)", "결론"], 1):
    c = ws.cell(hdr_row, j, h)
    c.fill, c.font, c.border = HDR, HDR_FONT, BORDER
cats = defaultdict(list)
for r in all_cases:
    cats[r["category"]].append(r)
row = hdr_row + 1
for cat, lst in cats.items():
    ex = [x for x in lst if is_exec(x)]
    p = sum(1 for x in ex if x["result"] == "PASS")
    avg = sum((x["match_pct"] or 0) for x in ex) / max(len(ex), 1)
    concl = "예상과 일치" if p == len(ex) and ex else ("일부 불일치 — 결함 시트 참고" if ex else "실행 불가")
    for j, v in enumerate([KR_CAT.get(cat, cat), len(lst), p, len(ex) - p, sum(1 for x in lst if x["result"] == "N/T"), round(avg, 1), concl], 1):
        c = ws.cell(row, j, v)
        c.border, c.font = BORDER, BASE_FONT
        if j == 7:
            c.fill = GREEN if concl == "예상과 일치" else YELLOW
    row += 1

# ------------------------------------------------------------- defects sheet
ws = sheet(wb, "결함·수정", ["ID", "심각도", "영역", "상태", "제목", "상세(재현/영향)", "조치", "관련 케이스", "재테스트 결과"], [8, 10, 16, 14, 40, 60, 50, 22, 26])
final_by_id = {r["id"]: r for r in all_cases}
for d in DEFECTS:
    ids = d[7]
    retest = ", ".join(f"{i}: {final_by_id[i]['result']} {final_by_id[i]['match_pct']}%" for i in ids if i in final_by_id) or "-"
    ws.append(list(d[:7]) + [", ".join(ids), retest])
style_rows(ws, result_col=4)

# --------------------------------------------------------- claims / limits
ws = sheet(wb, "웹↔앱 일치 점검", ["웹사이트 문구(수정 전)", "앱/서버 실제 동작", "판정", "조치(웹사이트)"], [46, 56, 12, 56])
for row_ in CLAIMS:
    ws.append(list(row_))
style_rows(ws, result_col=3)

ws = sheet(wb, "테스트 못한 것·한계", ["항목", "이유", "대체 검증/권장"], [40, 60, 60])
for row_ in LIMITS:
    ws.append(list(row_))
style_rows(ws)

# -------------------------------------------------------------- pytest sheet
if pytest_rows:
    ws = sheet(wb, "기존 pytest 회귀", ["모듈", "테스트", "결과", "시간(초)", "메시지"], [34, 70, 9, 10, 60])
    for row_ in pytest_rows:
        ws.append(list(row_))
    style_rows(ws, result_col=3)

out = DEST / (sys.argv[1] if len(sys.argv) > 1 else f"SooDaMate_QA_Report_{datetime.now():%Y-%m-%d}.xlsx")
wb.save(out)
print("saved", out)
print(f"executed={len(executed)} pass={n_pass} fail={n_fail} n/t={n_nt} case%={100.0 * n_pass / max(len(executed), 1):.1f} check%={100.0 * chk_pass / max(chk_total, 1):.1f} avg%={avg_match:.1f}")
print(f"baseline pass={base_pass}/{len(base_all)} avg={base_avg:.1f}")
